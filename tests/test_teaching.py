import json
from types import SimpleNamespace

import pytest

from jarvis.brain import TEACH_KICKOFF, Brain
from jarvis.config import load_config
from jarvis.knowledge import KnowledgeBase, chunk_text, fts_query
from jarvis.store import Store
from jarvis.tools.base import ToolContext, build_registry


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = load_config(overrides={"assistant": {"data_dir": str(tmp_path / "data")}})
    store = Store(cfg.db_path)
    kb = KnowledgeBase(store)
    return cfg, store, kb, tmp_path


def resp(content, stop="end_turn"):
    return SimpleNamespace(content=content, stop_reason=stop)


class FakeClient:
    def __init__(self, script):
        self.script, self.calls, self.messages = list(script), [], self

    def create(self, **kw):
        self.calls.append(json.loads(json.dumps(kw, default=str)))
        return self.script.pop(0)


def make_docs(folder):
    import docx
    import openpyxl
    from pptx import Presentation
    from pypdf import PdfWriter

    folder.mkdir()
    d = docx.Document()
    d.add_paragraph("Anchor camera spec: IP68 rated housing, 316L stainless, PoE+ 802.3at.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Lens", "2.8mm wide angle"
    d.save(folder / "anchor-camera.docx")
    wb = openpyxl.Workbook()
    wb.active.title = "Prices"
    wb.active.append(["Item", "Price"])
    wb.active.append(["Ubiquiti U7 Pro", 189])
    wb.save(folder / "pricing.xlsx")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Wave AVI pitch"
    s.placeholders[1].text = "Superyacht connectivity: Starlink plus 5G failover"
    prs.save(folder / "pitch.pptx")
    w = PdfWriter()
    w.add_blank_page(200, 200)
    w.write(folder / "blank-scan.pdf")
    (folder / "notes.md").write_text("# Site visit\nThe M/Y Aurora rack is in the lazarette.")
    (folder / "~$temp.docx").write_text("lock file")
    (folder / "image.png").write_bytes(b"\x89PNG")


def test_ingest_folder_all_formats_and_search(env):
    cfg, store, kb, tmp = env
    make_docs(tmp / "docs")
    res = kb.ingest_folder(tmp / "docs", topic="engineering")
    assert res["counts"]["added"] == 4
    assert any("blank-scan.pdf" in p for p in res["problems"])
    assert kb.search("what IP rating does the anchor camera have")[0]["title"] == "anchor-camera.docx"
    assert kb.search("U7 Pro price")[0]["title"] == "pricing.xlsx"
    assert kb.search("starlink failover")[0]["title"] == "pitch.pptx"
    assert kb.search("where is Aurora's rack", topic="engineering")[0]["title"] == "notes.md"
    assert kb.search("aurora", topic="other") == []
    # re-scan: unchanged; then edit + delete
    assert kb.ingest_folder(tmp / "docs")["counts"].get("unchanged") == 4
    (tmp / "docs" / "notes.md").write_text("# Site visit\nRack moved to the bridge deck.")
    (tmp / "docs" / "pitch.pptx").unlink()
    counts = kb.ingest_folder(tmp / "docs")["counts"]
    assert counts["updated"] == 1 and counts["removed"] == 1
    assert "bridge" in kb.search("rack location")[0]["text"]
    assert kb.search("starlink failover") == []


def test_notes_replace_by_title_and_read_source(env):
    _, _, kb, _ = env
    kb.add_note("Pricing policy", "Margin is 30 percent on hardware.", "business")
    kb.add_note("Pricing policy", "Margin is 35 percent on hardware, 50 on labour.", "business")
    notes = kb.list_sources(kind="note")
    assert len(notes) == 1
    hit = kb.search("hardware margin")[0]
    assert "35 percent" in hit["text"]
    long = "\n\n".join(f"Paragraph {i} " + "word " * 150 for i in range(30))
    sid = kb.add_note("Long manual", long)
    page = kb.read_source(sid)
    assert page["next_part"] and len(page["text"]) <= 8000
    page2 = kb.read_source(sid, page["next_part"])
    assert page2["parts_shown"].startswith(str(page["next_part"]))


def test_chunking_and_query_sanitising():
    chunks = chunk_text("x " * 5000)
    assert all(len(c) <= 1300 for c in chunks) and len(chunks) > 5
    assert fts_query('what is the "IP68" rating? AND OR') == '"ip68" OR "rating"'
    assert fts_query("the of a") == ""


def test_teaching_tools(env):
    cfg, store, kb, tmp = env
    reg = build_registry(cfg)
    ctx = ToolContext(cfg, store, kb=kb)
    out, err = reg.execute("learn_procedure", {
        "name": "Client quotation", "when_to_use": "When a client asks for a price",
        "steps": "1. Check pricing sheet\n2. Add 35% margin\n3. Send from sales@"}, ctx)
    assert not err and "saved" in out
    reg.execute("learn_procedure", {"name": "client quotation", "when_to_use": "Pricing requests",
                                    "steps": "1. Updated"}, ctx)
    assert len(store.list_procedures()) == 1
    got = json.loads(reg.execute("get_procedure", {"name": "Client Quotation"}, ctx)[0])
    assert got["steps"] == "1. Updated"
    out, _ = reg.execute("teach_note", {"title": "Wave entities", "content": "UK, NL, SG, TH, KH",
                                        "topic": "company"}, ctx)
    assert "Saved note" in out
    hits = json.loads(reg.execute("search_knowledge", {"query": "wave entities"}, ctx)[0])
    assert hits[0]["title"] == "Wave entities"
    (tmp / "a.md").write_text("Torque for M8 A4 bolts: 21 Nm")
    out = json.loads(reg.execute("learn_from_files", {"path": str(tmp / "a.md")}, ctx)[0])
    assert out["status"] == "added"
    stats = json.loads(reg.execute("list_knowledge", {"stats_only": True}, ctx)[0])
    assert stats["by_kind"]["note"]["sources"] == 1
    out, _ = reg.execute("remember", {"fact": "Never send quotes on Fridays", "category": "instruction"}, ctx)
    mid = int(out.split("#")[1].rstrip("."))
    reg.execute("update_memory", {"memory_id": mid, "fact": "Never send quotes after 5pm"}, ctx)
    assert store.memories_in(("instruction",))[0]["content"] == "Never send quotes after 5pm"


def test_system_prompt_has_teaching_context(env):
    cfg, store, kb, _ = env
    store.add_memory("Always answer in metric units", "instruction")
    store.add_memory("Wrong: Aurora is 60m. Correct: 72m", "correction")
    store.add_memory("Works with Starlink Maritime", "work")
    store.upsert_procedure("Weekly invoicing", "Every Friday", "1. ...")
    kb.add_note("Aurora refit", "The M/Y Aurora refit uses a Cisco core switch.", "clients")
    client = FakeClient([resp([{"type": "text", "text": "ok"}])])
    brain = Brain(cfg, store, build_registry(cfg), client=client, kb=kb)
    brain.run("which switch is on Aurora?")
    sys = client.calls[0]["system"]
    assert "Always answer in metric units" in sys.split("What you know")[0]
    assert "Correct: 72m" in sys
    assert "Weekly invoicing: Every Friday" in sys
    assert "Cisco core switch" in sys  # auto-recall
    assert "1 note" in sys and "clients" in sys
    assert "yachts and marine engineering" in sys and "TEACH MODE" not in sys
    names = {t.get("name") for t in client.calls[0]["tools"]}
    assert {"search_knowledge", "learn_procedure", "teach_note", "ask_expert"} <= names


def test_teach_mode_prompt(env):
    cfg, store, kb, _ = env
    client = FakeClient([resp([{"type": "text", "text": "What is your role?"}])])
    brain = Brain(cfg, store, build_registry(cfg), client=client, kb=kb)
    brain.mode = "teach"
    assert brain.run(TEACH_KICKOFF) == "What is your role?"
    assert "TEACH MODE IS ON" in client.calls[0]["system"]


def test_ask_expert_uses_expert_model(env):
    cfg, store, kb, _ = env
    client = FakeClient([
        resp([{"type": "tool_use", "id": "t1", "name": "ask_expert",
               "input": {"question": "Voltage drop on 30m of 6mm2 at 32A, 230V?", "context": "copper"}}], "tool_use"),
        resp([{"type": "server_tool_use", "id": "s", "name": "web_search", "input": {}}], "pause_turn"),
        resp([{"type": "text", "text": "About 4.9 V (2.1%)."}]),
        resp([{"type": "text", "text": "Roughly 4.9 volts, about 2 percent."}]),
    ])
    brain = Brain(cfg, store, build_registry(cfg), client=client, kb=kb)
    assert brain.run("voltage drop?", voice=True) == "Roughly 4.9 volts, about 2 percent."
    assert client.calls[1]["model"] == "claude-opus-5-5"
    assert "electrotechnics" in client.calls[1]["system"]
    assert "4.9 V" in client.calls[3]["messages"][2]["content"][0]["content"]


def test_backup_and_export(env):
    from jarvis.backup import backup

    cfg, store, kb, tmp = env
    store.add_memory("Prefers WhatsApp for urgent things", "preference")
    store.upsert_procedure("Commissioning", "New vessel network", "1. VLAN plan")
    kb.add_note("Suppliers", "Ubiquiti via distributor X")
    db, md = backup(store, kb, tmp / "bk", keep=2)
    text = md.read_text()
    assert "WhatsApp" in text and "VLAN plan" in text and "distributor X" in text
    import sqlite3

    assert sqlite3.connect(db).execute("SELECT COUNT(*) FROM memories").fetchone()[0] == 1


def test_scheduler_watch_folders(env):
    from jarvis.scheduler import JarvisScheduler

    cfg, store, kb, tmp = env
    (tmp / "w").mkdir()
    (tmp / "w" / "a.txt").write_text("Generator service interval 500 hours")
    cfg["knowledge"]["watch_folders"] = [{"path": str(tmp / "w"), "topic": "maintenance"}]
    brain = SimpleNamespace(kb=kb)
    sched = JarvisScheduler(cfg, store, brain, lambda *a: None)
    sched.sync_watch_folders()
    assert kb.search("generator service")[0]["topic"] == "maintenance"
    db, md = sched.run_backup()
    assert db.exists()


def test_cli_learn(env, capsys):
    from jarvis.__main__ import main

    cfg, _, _, tmp = env
    (tmp / "spec.md").write_text("Fibre backbone is OM4 multimode")
    main(["learn", str(tmp / "spec.md"), "--topic", "network"])
    out = capsys.readouterr().out
    assert "added" in out


def test_voice_teach_mode(env, monkeypatch):
    np = pytest.importorskip("numpy")
    import jarvis.voice.loop as loop

    cfg, store, kb, _ = env
    heard = iter(["let me teach you about my work", "I run Wave AVI", "stop teaching", "", "shut down jarvis"])
    rec_settings = []

    class FakeRecorder:
        def __init__(self, *a):
            self.silence_seconds, self.max_seconds = 1.2, 20

        def calibrate(self):
            pass

        def record(self, start_timeout=6):
            rec_settings.append((self.silence_seconds, start_timeout))
            return np.zeros(16000, dtype=np.int16)

    monkeypatch.setattr(loop, "Recorder", FakeRecorder)
    monkeypatch.setattr(loop, "Transcriber", lambda *a: SimpleNamespace(transcribe=lambda a: next(heard)))
    monkeypatch.setattr(loop, "WakeWord", lambda *a: SimpleNamespace(wait=lambda: None))
    monkeypatch.setattr(loop, "chime", lambda *a: None)
    client = FakeClient([
        resp([{"type": "text", "text": "I know nothing yet. What do you do?"}]),
        resp([{"type": "tool_use", "id": "t", "name": "remember",
               "input": {"fact": "Adrian runs Wave AVI", "category": "work"}}], "tool_use"),
        resp([{"type": "text", "text": "Noted. Which entities does Wave have?"}]),
        resp([{"type": "text", "text": "Today I learned you run Wave AVI."}]),
    ])
    brain = Brain(cfg, store, build_registry(cfg), client=client, kb=kb)
    spoken = []
    loop.VoiceAssistant(cfg, brain, SimpleNamespace(say=spoken.append)).run()
    assert spoken[1:4] == ["I know nothing yet. What do you do?", "Noted. Which entities does Wave have?",
                           "Today I learned you run Wave AVI."]
    assert store.search_memories("wave")[0]["content"] == "Adrian runs Wave AVI"
    assert brain.mode == "normal"
    assert rec_settings[1] == (2.2, 15.0)  # longer pauses + follow-up window while teaching
