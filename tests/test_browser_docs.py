import json
import threading
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

import pytest

from jarvis.config import load_config
from jarvis.store import Store
from jarvis.tools.base import ToolContext, build_registry

FORM = """<!doctype html><html><head><meta charset="utf-8"><title>Crew Visa Application</title></head><body>
<h1>Crew Visa Application</h1>
<form id="f" onsubmit="event.preventDefault(); document.getElementById('done').innerText='Application received for ' + document.getElementById('name').value;">
  <label for="name">Full name</label><input id="name" name="fullname" required>
  <label>Email <input type="email" name="email" placeholder="you@example.com"></label>
  <label for="country">Nationality</label>
  <select id="country" name="country"><option>Choose…</option><option>United Kingdom</option><option>Netherlands</option></select>
  <label><input type="checkbox" name="terms"> I accept the terms</label>
  <input type="file" name="passport" aria-label="Passport scan">
  <button type="submit">Submit application</button>
</form>
<p id="done"></p>
<div style="margin-top:2000px"><a href="#">Far away link</a></div>
</body></html>"""


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = load_config(overrides={"assistant": {"data_dir": str(tmp_path / "data")},
                                 "browser": {"headless": True}})
    store = Store(cfg.db_path)
    reg = build_registry(cfg)
    return cfg, store, reg, tmp_path


@pytest.fixture
def site(tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    (root / "form.html").write_text(FORM)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), partial(SimpleHTTPRequestHandler, directory=str(root)))
    srv.RequestHandlerClass.log_message = lambda *a: None
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}/form.html"
    srv.shutdown()


def ref_of(snapshot: str, needle: str) -> int:
    for line in snapshot.splitlines():
        if needle in line and line.startswith("["):
            return int(line[1:line.index("]")])
    raise AssertionError(f"{needle!r} not in snapshot:\n{snapshot}")


def test_browser_fills_and_submits_with_approval(env, site):
    cfg, store, reg, tmp = env
    asked = []
    ctx = ToolContext(cfg, store, confirm=lambda m: asked.append(m) or False)
    snap, err = reg.execute("browser_open", {"url": site}, ctx)
    assert not err, snap
    assert "Title: Crew Visa Application" in snap
    assert '"Full name"' in snap and "options=Choose…|United Kingdom|Netherlands" in snap
    assert "(scroll)" in snap.split("Far away link")[1].splitlines()[0]
    name, email = ref_of(snap, '"Full name"'), ref_of(snap, 'placeholder="you@example.com"')
    country, terms = ref_of(snap, "options="), ref_of(snap, "I accept the terms")
    passport, submit = ref_of(snap, "Passport scan"), ref_of(snap, "Submit application")
    (tmp / "passport.pdf").write_bytes(b"%PDF-1.4 test")
    out, err = reg.execute("browser_fill_form", {"fields": [
        {"ref": name, "value": "Adrian Example"}, {"ref": email, "value": "a@example.com"},
        {"ref": country, "value": "Netherlands"}, {"ref": terms, "value": True},
        {"ref": passport, "value": str(tmp / "passport.pdf")}]}, ctx)
    assert not err and out.count("✓") == 5, out
    assert 'value="Netherlands"' in out and "☑" in out
    # submit needs approval: declined first
    out, _ = reg.execute("browser_click", {"ref": submit}, ctx)
    assert "did not approve" in out and asked and "Submit application" in asked[0]
    ctx.confirm = lambda m: True
    out, _ = reg.execute("browser_click", {"ref": submit, "final_action": True}, ctx)
    assert "Application received for Adrian Example" in out
    shot, _ = reg.execute("browser_screenshot", {}, ctx)
    assert shot[1]["type"] == "image" and len(shot[1]["source"]["data"]) > 1000
    out, _ = reg.execute("browser_navigate", {"action": "list_tabs"}, ctx)
    assert "Crew Visa Application" in out
    out, err = reg.execute("browser_click", {"ref": 999}, ctx)
    assert err and "page changed" in out


def test_screenshot_reaches_claude_as_image(env, site):
    from types import SimpleNamespace

    from jarvis.brain import Brain

    cfg, store, reg, _ = env
    calls = []

    class Client:
        messages = None

        def __init__(self):
            self.messages = self
            self.script = [
                SimpleNamespace(stop_reason="tool_use", content=[
                    {"type": "tool_use", "id": "a", "name": "browser_open", "input": {"url": site}}]),
                SimpleNamespace(stop_reason="tool_use", content=[
                    {"type": "tool_use", "id": "b", "name": "browser_screenshot", "input": {}}]),
                SimpleNamespace(stop_reason="end_turn", content=[{"type": "text", "text": "It's a visa form."}]),
            ]

        def create(self, **kw):
            calls.append(json.loads(json.dumps(kw)))
            return self.script.pop(0)

    brain = Brain(cfg, store, reg, client=Client())
    assert brain.run("what's on that page?") == "It's a visa form."
    result = calls[2]["messages"][-1]["content"][0]
    assert result["content"][1]["type"] == "image"


def test_fill_word_document(env):
    import docx

    cfg, store, reg, tmp = env
    d = docx.Document()
    p = d.add_paragraph("Dear ")
    p.add_run("{{client")
    p.add_run("_name}},")  # placeholder split across runs
    d.add_paragraph("Quote valid until [DATE]. Vessel: ________")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Total"
    t.rows[0].cells[1].text = "{{total}}"
    d.sections[0].header.paragraphs[0].text = "Ref {{ref}}"
    src = tmp / "quote.docx"
    d.save(src)
    ctx = ToolContext(cfg, store)
    info = json.loads(reg.execute("read_document", {"path": str(src)}, ctx)[0])
    assert "[DATE]" in info["placeholders"] and "________" in info["placeholders"]
    out = json.loads(reg.execute("fill_word_document", {"path": str(src), "replacements": {
        "{{client_name}}": "Captain Smith", "[DATE]": "30 Nov 2026", "________": "M/Y Aurora",
        "{{total}}": "SGD 48,500", "{{ref}}": "WAVE-2041", "{{missing}}": "x"}}, ctx)[0])
    assert out["not_found"] == ["{{missing}}"]
    filled = docx.Document(out["saved"])
    body = "\n".join(p.text for p in filled.paragraphs)
    assert "Dear Captain Smith," in body and "30 Nov 2026" in body and "M/Y Aurora" in body
    assert filled.tables[0].rows[0].cells[1].text == "SGD 48,500"
    assert filled.sections[0].header.paragraphs[0].text == "Ref WAVE-2041"
    assert "{{client" in "\n".join(p.text for p in docx.Document(src).paragraphs)  # original untouched


def test_create_word_document(env):
    import docx

    cfg, store, reg, tmp = env
    out, _ = reg.execute("create_word_document", {"output_path": str(tmp / "letter"), "content":
                         "# Proposal\n\nDear **Captain**,\n\n- Starlink\n- 5G failover\n1. Survey"},
                         ToolContext(cfg, store))
    d = docx.Document(tmp / "letter.docx")
    assert [p.style.name for p in d.paragraphs] == ["Heading 1", "Normal", "List Bullet", "List Bullet", "List Number"]
    assert d.paragraphs[1].runs[1].bold


def test_fill_pdf_form(env):
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    cfg, store, reg, tmp = env
    src = tmp / "form.pdf"
    c = canvas.Canvas(str(src))
    c.drawString(50, 750, "Marina berth application")
    c.acroForm.textfield(name="vessel_name", x=50, y=700, width=200, height=20)
    c.acroForm.textfield(name="loa_m", x=50, y=660, width=100, height=20)
    c.acroForm.checkbox(name="insurance", x=50, y=620, size=15)
    c.acroForm.choice(name="season", x=50, y=580, width=120, height=20, options=["Summer", "Winter"], value="Summer")
    c.save()
    ctx = ToolContext(cfg, store)
    fields = json.loads(reg.execute("pdf_form_fields", {"path": str(src)}, ctx)[0])
    assert {f["name"] for f in fields} == {"vessel_name", "loa_m", "insurance", "season"}
    out = json.loads(reg.execute("fill_pdf_form", {"path": str(src), "values": {
        "vessel_name": "Aurora", "loa_m": "72", "insurance": True, "season": "Winter", "bogus": 1}}, ctx)[0])
    assert out["unknown_fields"] == ["bogus"]
    got = {k: str(v.get("/V")) for k, v in PdfReader(out["saved"]).get_fields().items()}
    assert got["vessel_name"] == "Aurora" and got["loa_m"] == "72" and got["season"] == "Winter"
    assert got["insurance"] not in ("/Off", "None")


def test_fill_excel(env):
    import openpyxl

    cfg, store, reg, tmp = env
    wb = openpyxl.Workbook()
    wb.active.title = "Quote"
    wb.active["A1"] = "Item"
    wb.save(tmp / "q.xlsx")
    out = json.loads(reg.execute("fill_excel", {"path": str(tmp / "q.xlsx"), "cells": {
        "Quote!B2": 189, "B3": 2, "Quote!B4": "=B2*B3"}}, ToolContext(cfg, store))[0])
    ws = openpyxl.load_workbook(out["saved"])["Quote"]
    assert ws["B2"].value == 189 and ws["B4"].value == "=B2*B3"
    out, _ = reg.execute("fill_excel", {"path": str(tmp / "q.xlsx"), "cells": {"Nope!A1": 1}}, ToolContext(cfg, store))
    assert "No sheet" in out


def test_child_has_no_browser_or_document_tools(env):
    from jarvis.profiles import Profile, profile_config

    cfg, *_ = env
    kid = build_registry(profile_config(cfg, Profile("son", "Max", "child", 12)))
    assert not [n for n in kid.names() if n.startswith(("browser_", "fill_", "pdf_", "create_word", "read_document"))]
