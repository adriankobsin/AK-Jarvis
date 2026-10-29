"""Knowledge base: everything you teach Jarvis — documents, folders, web pages and notes —
chunked and indexed with SQLite FTS5 (BM25 ranking). Fully local, no embedding service."""
from __future__ import annotations

import email
import email.policy
import hashlib
import logging
import os
import re
from datetime import datetime
from pathlib import Path

log = logging.getLogger("jarvis.knowledge")

SCHEMA = """
CREATE TABLE IF NOT EXISTS kb_sources (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    kind TEXT NOT NULL,            -- file | url | note
    uri TEXT NOT NULL UNIQUE,
    title TEXT NOT NULL,
    topic TEXT DEFAULT '',
    hash TEXT,
    mtime REAL,
    chunks INTEGER DEFAULT 0,
    chars INTEGER DEFAULT 0,
    added_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE VIRTUAL TABLE IF NOT EXISTS kb_chunks USING fts5(
    content, title, topic, source_id UNINDEXED, seq UNINDEXED,
    tokenize = 'porter unicode61'
);
"""

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml", ".log",
            ".ini", ".cfg", ".conf", ".xml", ".rtf", ".py", ".js", ".ts", ".sql", ".sh"}
DOC_EXT = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".html", ".htm", ".eml"}
SUPPORTED = TEXT_EXT | DOC_EXT
STOPWORDS = set("""a an and are as at be but by can could did do does for from had has have how
i if in into is it its me my of on or our so than that the their them then there these they
this to was we were what when where which who why will with would you your about please
tell show give know""".split())


# ------------------------------------------------------------------ extraction
def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        return path.read_text(errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join(f"[Page {i + 1}]\n{p.extract_text() or ''}" for i, p in enumerate(reader.pages))
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"## Sheet: {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 5000:
                    parts.append("…(truncated)")
                    break
                vals = ["" if v is None else str(v) for v in row]
                if any(vals):
                    parts.append(" | ".join(vals).rstrip(" |"))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
            if slide.has_notes_slide:
                texts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
            parts.append(f"[Slide {i}]\n" + "\n".join(t for t in texts if t.strip()))
        return "\n\n".join(parts)
    if ext in (".html", ".htm"):
        return html_to_text(path.read_text(errors="replace"))
    if ext == ".eml":
        msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
        body = msg.get_body(preferencelist=("plain", "html"))
        content = body.get_content() if body else ""
        if body and body.get_content_subtype() == "html":
            content = html_to_text(content)
        head = "\n".join(f"{h}: {msg.get(h, '')}" for h in ("From", "To", "Cc", "Date", "Subject"))
        return f"{head}\n\n{content}"
    raise ValueError(f"Unsupported file type: {ext}")


def html_to_text(html: str) -> str:
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
            tag.decompose()
        text = soup.get_text("\n")
    except ImportError:
        text = re.sub(r"<[^>]+>", " ", html)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        return []
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    cur = ""
    for p in paras:
        while len(p) > size:  # very long paragraph: hard split on sentence/space
            cut = max(p.rfind(". ", 0, size), p.rfind(" ", 0, size))
            cut = cut + 1 if cut > size // 2 else size
            pieces = p[:cut].strip()
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(pieces)
            p = p[max(0, cut - overlap):].strip()
        if len(cur) + len(p) + 2 <= size:
            cur = f"{cur}\n\n{p}" if cur else p
        else:
            if cur:
                chunks.append(cur)
            tail = cur[-overlap:] if cur else ""
            tail = tail[tail.find(" ") + 1:] if " " in tail else tail
            cur = f"{tail}\n\n{p}" if tail else p
    if cur:
        chunks.append(cur)
    return chunks


def fts_query(text: str) -> str:
    words = re.findall(r"[\w][\w\-\.]*", text.lower())
    words = [w.strip(".-") for w in words]
    words = [w for w in words if w and w not in STOPWORDS and (len(w) > 1 or w.isdigit())]
    seen: list[str] = []
    for w in words:
        if w not in seen:
            seen.append(w)
    return " OR ".join('"' + w.replace('"', "") + '"' for w in seen[:24])


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:80] or "note"


# ------------------------------------------------------------------ knowledge base
class KnowledgeBase:
    def __init__(self, store, max_file_mb: float = 25):
        self.store = store
        self._db = store._db
        self._lock = store._lock
        self.max_bytes = int(max_file_mb * 2**20)
        with self._lock:
            self._db.executescript(SCHEMA)
            self._db.commit()

    # -- writing ---------------------------------------------------------
    def add_text(self, uri: str, title: str, text: str, kind: str, topic: str = "",
                 hash_: str | None = None, mtime: float | None = None) -> int:
        chunks = chunk_text(text)
        ts = datetime.now().replace(microsecond=0).isoformat()
        with self._lock:
            row = self._db.execute("SELECT id FROM kb_sources WHERE uri=?", (uri,)).fetchone()
            if row:
                sid = row[0]
                self._db.execute("DELETE FROM kb_chunks WHERE source_id=?", (sid,))
                self._db.execute(
                    "UPDATE kb_sources SET title=?, topic=?, hash=?, mtime=?, chunks=?, chars=?, updated_at=? WHERE id=?",
                    (title, topic, hash_, mtime, len(chunks), len(text), ts, sid))
            else:
                sid = self._db.execute(
                    "INSERT INTO kb_sources(kind, uri, title, topic, hash, mtime, chunks, chars, added_at, updated_at)"
                    " VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (kind, uri, title, topic, hash_, mtime, len(chunks), len(text), ts, ts)).lastrowid
            self._db.executemany(
                "INSERT INTO kb_chunks(content, title, topic, source_id, seq) VALUES (?,?,?,?,?)",
                [(c, title, topic, sid, i) for i, c in enumerate(chunks)])
            self._db.commit()
        return sid

    def add_note(self, title: str, content: str, topic: str = "") -> int:
        return self.add_text(f"note:{slug(title)}", title, f"{title}\n\n{content}", "note", topic)

    def ingest_file(self, path: str | Path, topic: str = "", force: bool = False) -> str:
        """Returns 'added' | 'updated' | 'unchanged' | 'skipped: reason'."""
        p = Path(os.path.expanduser(str(path))).resolve()
        if not p.is_file():
            return "skipped: not a file"
        if p.suffix.lower() not in SUPPORTED:
            return "skipped: unsupported type"
        if p.name.startswith(("~$", ".")):
            return "skipped: temp/hidden"
        st = p.stat()
        if st.st_size > self.max_bytes:
            return "skipped: too large"
        uri = p.as_uri()
        with self._lock:
            row = self._db.execute("SELECT mtime, hash FROM kb_sources WHERE uri=?", (uri,)).fetchone()
        if row and not force and row[0] == st.st_mtime:
            return "unchanged"
        data = p.read_bytes()
        h = hashlib.sha1(data).hexdigest()
        if row and not force and row[1] == h:
            with self._lock:
                self._db.execute("UPDATE kb_sources SET mtime=? WHERE uri=?", (st.st_mtime, uri))
                self._db.commit()
            return "unchanged"
        try:
            text = extract_text(p)
        except ImportError as e:
            return f"skipped: missing parser ({e.name})"
        except Exception as e:
            return f"skipped: {type(e).__name__}: {e}"
        if not re.sub(r"\[(Page|Slide) \d+\]|## Sheet: .*", "", text).strip():
            return "skipped: no text (scanned image?)"
        self.add_text(uri, p.name, f"{p.name}\n{text}", "file", topic, h, st.st_mtime)
        return "updated" if row else "added"

    def ingest_folder(self, folder: str | Path, topic: str = "", recursive: bool = True) -> dict:
        root = Path(os.path.expanduser(str(folder))).resolve()
        if not root.is_dir():
            raise FileNotFoundError(f"No such folder: {root}")
        counts: dict[str, int] = {}
        errors: list[str] = []
        seen: set[str] = set()
        it = root.rglob("*") if recursive else root.glob("*")
        for f in it:
            if any(part.startswith(".") for part in f.relative_to(root).parts):
                continue
            if not f.is_file() or f.suffix.lower() not in SUPPORTED:
                continue
            seen.add(f.resolve().as_uri())
            status = self.ingest_file(f, topic)
            key = status.split(":")[0]
            counts[key] = counts.get(key, 0) + 1
            if status.startswith("skipped") and "temp" not in status:
                errors.append(f"{f.name}: {status}")
        # remove sources for files that were deleted from this folder
        prefix = root.as_uri().rstrip("/") + "/"
        removed = 0
        for s in self.list_sources(kind="file"):
            if s["uri"].startswith(prefix) and s["uri"] not in seen:
                self.delete_source(s["id"])
                removed += 1
        if removed:
            counts["removed"] = removed
        return {"folder": str(root), "counts": counts, "problems": errors[:20]}

    def ingest_url(self, url: str, topic: str = "") -> str:
        import requests

        r = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0 Jarvis"})
        r.raise_for_status()
        ctype = r.headers.get("content-type", "")
        if "pdf" in ctype:
            import io

            from pypdf import PdfReader

            text = "\n\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(r.content)).pages)
            title = url.rsplit("/", 1)[-1]
        else:
            m = re.search(r"<title[^>]*>(.*?)</title>", r.text, re.I | re.S)
            title = re.sub(r"\s+", " ", m.group(1)).strip() if m else url
            text = html_to_text(r.text)
        self.add_text(url, title, f"{title}\n{url}\n\n{text}", "url", topic)
        return title

    def delete_source(self, id_or_uri) -> bool:
        with self._lock:
            row = self._db.execute("SELECT id FROM kb_sources WHERE id=? OR uri=?",
                                   (id_or_uri, str(id_or_uri))).fetchone()
            if not row:
                return False
            self._db.execute("DELETE FROM kb_chunks WHERE source_id=?", (row[0],))
            self._db.execute("DELETE FROM kb_sources WHERE id=?", (row[0],))
            self._db.commit()
            return True

    # -- reading ---------------------------------------------------------
    def search(self, query: str, limit: int = 6, topic: str = "") -> list[dict]:
        q = fts_query(query)
        if not q:
            return []
        sql = ("SELECT c.content, c.source_id, c.seq, s.title, s.uri, s.kind, s.topic, bm25(kb_chunks, 1.0, 3.0, 2.0) AS score "
               "FROM kb_chunks c JOIN kb_sources s ON s.id = c.source_id WHERE kb_chunks MATCH ?")
        args: list = [q]
        if topic:
            sql += " AND lower(s.topic) = ?"
            args.append(topic.lower())
        sql += " ORDER BY score LIMIT ?"
        args.append(limit * 3)
        with self._lock:
            rows = self._db.execute(sql, args).fetchall()
        out, per_source = [], {}
        for r in rows:  # diversify: max 2 chunks per source
            per_source[r[1]] = per_source.get(r[1], 0) + 1
            if per_source[r[1]] > 2:
                continue
            out.append({"source_id": r[1], "title": r[3], "uri": r[4], "kind": r[5], "topic": r[6],
                        "part": r[2], "score": round(-r[7], 2), "text": r[0]})
            if len(out) >= limit:
                break
        return out

    def read_source(self, id_or_uri, start_part: int = 0, max_chars: int = 8000) -> dict | None:
        with self._lock:
            src = self._db.execute("SELECT id, title, uri, chunks FROM kb_sources WHERE id=? OR uri=?",
                                   (id_or_uri, str(id_or_uri))).fetchone()
            if not src:
                return None
            rows = self._db.execute(
                "SELECT seq, content FROM kb_chunks WHERE source_id=? AND CAST(seq AS INTEGER)>=? "
                "ORDER BY CAST(seq AS INTEGER)", (src[0], int(start_part))).fetchall()
        text, last = "", int(start_part) - 1
        for seq, content in rows:
            if text and len(text) + len(content) > max_chars:
                break
            text += ("\n\n" if text else "") + content
            last = int(seq)
        return {"source_id": src[0], "title": src[1], "uri": src[2], "total_parts": src[3],
                "parts_shown": f"{start_part}-{last}",
                "next_part": last + 1 if last + 1 < src[3] else None, "text": text}

    def list_sources(self, kind: str = "", topic: str = "") -> list[dict]:
        sql, args = "SELECT * FROM kb_sources WHERE 1=1", []
        if kind:
            sql += " AND kind=?"
            args.append(kind)
        if topic:
            sql += " AND lower(topic)=?"
            args.append(topic.lower())
        with self._lock:
            cur = self._db.execute(sql + " ORDER BY updated_at DESC", args)
            cols = [c[0] for c in cur.description]
            return [dict(zip(cols, r)) for r in cur.fetchall()]

    def stats(self) -> dict:
        with self._lock:
            rows = self._db.execute(
                "SELECT kind, COUNT(*), COALESCE(SUM(chunks),0), COALESCE(SUM(chars),0) FROM kb_sources GROUP BY kind").fetchall()
            topics = [r[0] for r in self._db.execute(
                "SELECT DISTINCT topic FROM kb_sources WHERE topic!='' ORDER BY topic").fetchall()]
        return {"by_kind": {k: {"sources": n, "chunks": c, "chars": ch} for k, n, c, ch in rows},
                "topics": topics}
