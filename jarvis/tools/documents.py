"""Read and fill in documents: Word (.docx), PDF forms, Excel; create new Word documents.
Originals are never overwritten unless explicitly requested — output defaults to '<name>-filled'."""
from __future__ import annotations

import re
from pathlib import Path

PLACEHOLDER = re.compile(r"\{\{[^}]+\}\}|\[[A-Z][A-Za-z0-9 _/\-]{1,40}\]|_{4,}|<<[^>]+>>")


def _out_path(src: Path, output_path: str | None) -> Path:
    if output_path:
        return Path(output_path).expanduser()
    return src.with_name(f"{src.stem}-filled{src.suffix}")


def _iter_paragraphs(doc):
    def from_container(c):
        for p in c.paragraphs:
            yield p
        for t in getattr(c, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from from_container(cell)

    yield from from_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            try:
                yield from from_container(part)
            except Exception:
                pass


def replace_in_paragraph(p, old: str, new: str) -> int:
    if old not in p.text:
        return 0
    count = 0
    for run in p.runs:  # keep formatting when the text sits inside one run
        if old in run.text:
            count += run.text.count(old)
            run.text = run.text.replace(old, new)
    if old in p.text and p.runs:  # spans several runs: merge into the first run
        full = "".join(r.text for r in p.runs)
        count += full.count(old)
        p.runs[0].text = full.replace(old, new)
        for r in p.runs[1:]:
            r.text = ""
    return count


def _record(ctx, path) -> None:
    sink = ctx.extras.get("media")
    if sink:
        try:
            sink({"source": str(Path(path).resolve()), "title": Path(path).name, "kind": "file", "auto": False})
        except Exception:
            pass


def register(reg, config) -> None:
    @reg.tool(
        "read_document",
        "Read the text of a document on this computer (Word, PDF, Excel, PowerPoint, text, email) "
        "and list any placeholders/blanks found, e.g. {{name}}, [DATE], ____ . For PDFs with "
        "form fields use pdf_form_fields instead.",
        {"path": {"type": "string"}, "max_chars": {"type": "integer"}},
        ["path"],
    )
    def read_document(ctx, path: str, max_chars: int = 15000):
        from ..knowledge import extract_text

        p = Path(path).expanduser()
        if not p.is_file():
            return f"No such file: {p}"
        text = extract_text(p)
        blanks = sorted(set(PLACEHOLDER.findall(text)))[:60]
        return {"file": str(p), "chars": len(text), "placeholders": blanks, "text": text[:int(max_chars)]}

    @reg.tool(
        "fill_word_document",
        "Fill in a Word (.docx) document by replacing text: placeholders like {{client_name}} or "
        "[DATE], underscores ____, or any exact text. Works in body, tables, headers and footers. "
        "Saves to output_path (default '<name>-filled.docx'); the original is untouched.",
        {
            "path": {"type": "string"},
            "replacements": {"type": "object", "description": "Map of exact existing text → new text"},
            "output_path": {"type": "string"},
        },
        ["path", "replacements"],
    )
    def fill_word_document(ctx, path: str, replacements: dict, output_path: str = ""):
        import docx

        src = Path(path).expanduser()
        doc = docx.Document(str(src))
        counts = {k: 0 for k in replacements}
        for p in _iter_paragraphs(doc):
            for old, new in replacements.items():
                counts[old] += replace_in_paragraph(p, old, str(new))
        out = _out_path(src, output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        missing = [k for k, v in counts.items() if v == 0]
        _record(ctx, out)
        return {"saved": str(out), "replaced": {k: v for k, v in counts.items() if v},
                "not_found": missing}

    @reg.tool(
        "create_word_document",
        "Create a new Word document from simple markdown: '# ' / '## ' headings, '- ' bullets, "
        "'1. ' numbered items, **bold** text, blank lines between paragraphs. For letters, "
        "proposals, reports, applications, cover letters.",
        {"output_path": {"type": "string"}, "content": {"type": "string"}},
        ["output_path", "content"],
    )
    def create_word_document(ctx, output_path: str, content: str):
        import docx

        doc = docx.Document()
        for block in content.split("\n"):
            line = block.rstrip()
            if not line.strip():
                continue
            m = re.match(r"^(#{1,3})\s+(.*)", line)
            if m:
                doc.add_heading(m.group(2), level=len(m.group(1)))
                continue
            style = None
            if re.match(r"^\s*[-*•]\s+", line):
                style, line = "List Bullet", re.sub(r"^\s*[-*•]\s+", "", line)
            elif re.match(r"^\s*\d+[.)]\s+", line):
                style, line = "List Number", re.sub(r"^\s*\d+[.)]\s+", "", line)
            para = doc.add_paragraph(style=style)
            for i, part in enumerate(re.split(r"\*\*", line)):
                if part:
                    para.add_run(part).bold = (i % 2 == 1)
        out = Path(output_path).expanduser()
        if out.suffix.lower() != ".docx":
            out = out.with_suffix(".docx")
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        _record(ctx, out)
        return f"Created {out}"

    @reg.tool(
        "pdf_form_fields",
        "List the fillable fields of a PDF form: names, types, current values and choices.",
        {"path": {"type": "string"}},
        ["path"],
    )
    def pdf_form_fields(ctx, path: str):
        from pypdf import PdfReader

        reader = PdfReader(str(Path(path).expanduser()))
        fields = reader.get_fields() or {}
        if not fields:
            return "This PDF has no fillable form fields (it may be a flat/scanned form)."
        out = []
        for name, f in fields.items():
            ft = f.get("/FT")
            item = {"name": name, "type": {"/Tx": "text", "/Btn": "checkbox/radio", "/Ch": "choice",
                                           "/Sig": "signature"}.get(ft, str(ft)),
                    "value": str(f.get("/V", "") or "")}
            if f.get("/_States_"):
                item["states"] = [str(s) for s in f["/_States_"]]
            if f.get("/Opt"):
                item["choices"] = [str(o if not isinstance(o, list) else o[-1]) for o in f["/Opt"]]
            if f.get("/TU"):
                item["tooltip"] = str(f["/TU"])
            out.append(item)
        return out

    @reg.tool(
        "fill_pdf_form",
        "Fill a PDF form's fields (names from pdf_form_fields). Checkboxes accept true/false. "
        "Saves to output_path (default '<name>-filled.pdf'); the original is untouched.",
        {
            "path": {"type": "string"},
            "values": {"type": "object", "description": "field name → value"},
            "output_path": {"type": "string"},
        },
        ["path", "values"],
    )
    def fill_pdf_form(ctx, path: str, values: dict, output_path: str = ""):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import NameObject

        src = Path(path).expanduser()
        reader = PdfReader(str(src))
        fields = reader.get_fields() or {}
        writer = PdfWriter()
        writer.append(reader)
        resolved, unknown = {}, []
        for k, v in values.items():
            if k not in fields:
                unknown.append(k)
                continue
            f = fields[k]
            if f.get("/FT") == "/Btn" and isinstance(v, (bool, str)) and str(v).lower() in (
                    "true", "false", "yes", "no", "on", "off", "1", "0"):
                on = next((str(s) for s in f.get("/_States_", []) if str(s) != "/Off"), "/Yes")
                resolved[k] = on if str(v).lower() in ("true", "yes", "on", "1") else "/Off"
            else:
                resolved[k] = str(v)
        for page in writer.pages:
            writer.update_page_form_field_values(page, resolved, auto_regenerate=False)
        writer.set_need_appearances_writer(True)
        # make checkbox appearance states follow their values
        for page in writer.pages:
            for annot in page.get("/Annots", []) or []:
                a = annot.get_object()
                name = a.get("/T")
                if name in resolved and a.get("/FT", a.get("/Parent", {}).get("/FT") if a.get("/Parent") else None) == "/Btn":
                    a[NameObject("/AS")] = NameObject(resolved[name])
        out = _out_path(src, output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        with open(out, "wb") as fh:
            writer.write(fh)
        _record(ctx, out)
        return {"saved": str(out), "filled": list(resolved), "unknown_fields": unknown}

    @reg.tool(
        "fill_excel",
        "Write values into an Excel workbook: cells as {'Sheet1!B3': value} (or 'B3' for the first "
        "sheet). Formulas allowed as strings starting with '='. Saves to output_path (default "
        "'<name>-filled.xlsx'); the original is untouched.",
        {"path": {"type": "string"}, "cells": {"type": "object"}, "output_path": {"type": "string"}},
        ["path", "cells"],
    )
    def fill_excel(ctx, path: str, cells: dict, output_path: str = ""):
        import openpyxl

        src = Path(path).expanduser()
        wb = openpyxl.load_workbook(str(src), keep_vba=src.suffix.lower() == ".xlsm")
        written = []
        for ref, value in cells.items():
            sheet, cell = ref.rsplit("!", 1) if "!" in ref else (wb.sheetnames[0], ref)
            sheet = sheet.strip("'")
            if sheet not in wb.sheetnames:
                return f"No sheet named {sheet!r}. Sheets: {wb.sheetnames}"
            wb[sheet][cell] = value
            written.append(f"{sheet}!{cell}")
        out = _out_path(src, output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        _record(ctx, out)
        return {"saved": str(out), "cells_written": written}
